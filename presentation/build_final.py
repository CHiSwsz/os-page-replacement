#!/usr/bin/env python3
"""Build the final contest deck (Chinese, 5 algorithms, our algorithm = 'Adaptive').

Reads results/v5/*.csv (internal directory name; never shown on slides).
Parameter study: presentation/param_sweep.csv (produced with the contest binary).
"""
import csv, json, statistics
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / 'results/v5'
FONT = 'Microsoft YaHei'
PURPLE, INK, GRAY, TEAL, LIGHT, AMBER, BLUE, SLATE = '6536A0', '222638', '637083', '087F78', 'F3EFF9', 'D98A0B', '2F6FD0', '334155'
COL = {'fifo': '9AA3B2', 'lru': AMBER, 'second_chance': BLUE, 'opt': SLATE, 'ours': TEAL, 'lfu': 'C9CED8', 'random': 'C9CED8'}
NAME = {'fifo': 'FIFO', 'lru': 'LRU', 'second_chance': 'Second chance', 'opt': 'MIN (OPT)', 'ours': '我们的算法', 'lfu': 'LFU', 'random': 'Random'}
FIVE = ['fifo', 'lru', 'second_chance', 'opt', 'ours']
FRAMES = [5, 10, 20, 50, 100]
TRACES = [('qsort', 'qsort'), ('lu', 'lu'), ('mm16', 'mm16'), ('mm32', 'mm32')]

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
notes = []


def rgb(h): return RGBColor.from_string(h)


def read(name):
    with (RES / (name + '.csv')).open() as f: return list(csv.DictReader(f))


def key(a): return 'adaptive_v5' if a == 'ours' else a


def tb(s, text, x, y, w, h, size=20, color=INK, bold=False, align=None, anchor=None, space=6):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.05); tf.margin_top = tf.margin_bottom = Inches(.03)
    if anchor: tf.vertical_anchor = anchor
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
        p.space_after = Pt(space)
        if align: p.alignment = align
    return shape


def box(s, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=.08):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = rgb(fill)
    if line: b.line.color.rgb = rgb(line); b.line.width = Pt(1.5)
    else: b.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0] = radius
    b.shadow.inherit = False
    return b


def card(s, x, y, w, h, title, body, head=PURPLE, fill=LIGHT, tsize=20, bsize=15, line=None):
    box(s, x, y, w, h, fill, line)
    hb = box(s, x, y, w, .55, head, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=.3)
    tb(s, title, x, y + .02, w, .5, tsize, 'FFFFFF', True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, body, x + .1, y + .68, w - .2, h - .75, bsize, INK, space=4)


def slide(title, script, seconds):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.333, .12, PURPLE, shape=MSO_SHAPE.RECTANGLE)
    tb(s, title, .55, .3, 12.2, .75, 30, PURPLE, True)
    box(s, 0, 7.12, 13.333, .38, 'F0EAF8', shape=MSO_SHAPE.RECTANGLE)
    tb(s, '页面替换算法设计  ·  Demand paging', .55, 7.16, 8, .3, 11, GRAY)
    tb(s, str(len(prs.slides)), 12.2, 7.16, .6, .3, 11, GRAY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = (f'建议用时：{seconds} 秒\n' if seconds else '附录，不讲\n') + script
    notes.append((title, seconds, script))
    return s


def table(s, headers, rows, x=.6, y=1.5, w=12.1, size=16, hl_col=None, rowh=.55, colw=None):
    shape = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(rowh * (len(rows) + 1)))
    t = shape.table
    if colw:
        for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate([headers, *rows]):
        for j, txt in enumerate(row):
            c = t.cell(i, j); c.text = str(txt); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(PURPLE if i == 0 else ('E3F4F2' if hl_col == j else ('F6F3FA' if i % 2 else 'FFFFFF')))
            for p in c.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.name = FONT; r.font.size = Pt(size)
                    r.font.bold = (i == 0 or hl_col == j)
                    r.font.color.rgb = rgb('FFFFFF' if i == 0 else (TEAL if hl_col == j else INK))
    return shape


def style_axis(c, size=12):
    c.category_axis.tick_labels.font.size = Pt(size); c.category_axis.tick_labels.font.name = FONT
    c.value_axis.tick_labels.font.size = Pt(size - 1); c.value_axis.tick_labels.font.name = FONT
    c.value_axis.minimum_scale = 0
    c.value_axis.major_gridlines.format.line.color.rgb = rgb('E5E7EB')
    c.value_axis.format.line.fill.background()


def line_chart(s, cats, series, x, y, w, h, title=None, legend=True):
    data = CategoryChartData(); data.categories = cats
    for name, vals, _ in series: data.add_series(name, vals)
    c = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    c.has_legend = legend
    if legend:
        c.legend.position = XL_LEGEND_POSITION.BOTTOM; c.legend.include_in_layout = False
        c.legend.font.size = Pt(12); c.legend.font.name = FONT
    if title:
        c.has_title = True; c.chart_title.text_frame.text = title
        r = c.chart_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(15); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = rgb(INK)
    else: c.has_title = False
    style_axis(c)
    for ser, (_, _, k) in zip(c.series, series):
        ser.smooth = False
        ser.format.line.color.rgb = rgb(COL[k]); ser.format.line.width = Pt(4 if k == 'ours' else 2.25)
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(COL[k])
        ser.marker.format.line.color.rgb = rgb(COL[k]); ser.marker.size = 8 if k == 'ours' else 6
        if k == 'opt': ser.format.line.dash_style = 4  # dash
    return c


def bar_chart(s, cats, vals, colors, x, y, w, h, fmt='#,##0', horizontal=False, size=13):
    data = CategoryChartData(); data.categories = cats; data.add_series('s', vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    c = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    c.has_legend = False; c.has_title = False; style_axis(c, size)
    if horizontal: c.category_axis.reverse_order = True
    c.value_axis.visible = False; c.value_axis.has_major_gridlines = False
    pl = c.plots[0]; pl.gap_width = 45; pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = fmt; dl.number_format_is_linked = False
    dl.font.size = Pt(size); dl.font.name = FONT; dl.font.bold = True; dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for i, col in enumerate(colors):
        pt = c.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(col)
    return c


def main():
    course = read('course_summary'); means = read('course_means')
    val = lambda d, f, a: float(next(r['faults_mean'] for r in course if r['dataset'] == d and int(r['frames']) == f and r['algorithm'] == key(a)))
    tmean = lambda d, a: statistics.mean(val(d, f, a) for f in FRAMES)
    overall = lambda a: statistics.mean(tmean(d, a) for d, _ in TRACES)
    O = {a: overall(a) for a in FIVE + ['lfu', 'random']}
    vs_lru = 100 * (1 - O['ours'] / O['lru']); vs_opt = 100 * (O['ours'] / O['opt'] - 1)
    vs_best_classic = 100 * (1 - O['ours'] / min(O['fifo'], O['lru'], O['second_chance']))
    perf = json.loads((RES / 'performance.json').read_text())
    mm = [val('mm32', f, 'ours') for f in FRAMES]

    # 1 title ---------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.333, 7.5, PURPLE, shape=MSO_SHAPE.RECTANGLE)
    box(s, .8, 3.55, 1.6, .07, 'FFFFFF', shape=MSO_SHAPE.RECTANGLE)
    tb(s, '页面替换算法设计', .8, 1.6, 11.5, 1.1, 54, 'FFFFFF', True)
    tb(s, 'Demand paging 下的自适应预测式页面替换', .8, 2.7, 11.5, .7, 28, 'E9DDF7')
    tb(s, 'One-Day Hackathon 汇报  ·  10.29', .8, 3.85, 11.5, .5, 20, 'E9DDF7')
    s.notes_slide.notes_text_frame.text = '建议用时：10 秒\n开场：我们实现了 FIFO、LRU、Second chance、MIN 四个经典算法，并设计了自己的自适应算法。'
    notes.append(('页面替换算法设计', 10, '开场：我们实现了 FIFO、LRU、Second chance、MIN 四个经典算法，并设计了自己的自适应算法。'))

    # 2 problem -------------------------------------------------------------
    s = slide('目标：用最少的缺页，处理没见过的访问序列',
              '题目是 demand paging：进程启动时没有页面，缺页且内存满时必须选一个受害页换出，不能预取。现场评分是在没见过的 trace 上，页框 5、10、20、50、100 五档的平均缺页次数，越少越好。关键矛盾：最优算法 MIN 需要知道未来，真实系统不可能；所以我们的思路是从已经发生的历史里学习，去估计未来。',
              40)
    card(s, .6, 1.5, 3.9, 2.5, '规则', '• 只做 demand paging，不预取\n• 每次缺页只换入当前页\n• 页框 5/10/20/50/100\n• 每档从空内存开始', bsize=16)
    card(s, 4.7, 1.5, 3.9, 2.5, '评分', '• 现场未知 trace\n• 五档缺页次数的算术平均\n• 越少越好', bsize=17)
    card(s, 8.8, 1.5, 3.9, 2.5, '难点', '• MIN 要读未来，现实做不到\n• 经典算法只看一种信号\n• 访问模式因程序而异', bsize=17, head=TEAL)
    box(s, .6, 4.45, 12.1, 1.9, 'E3F4F2')
    tb(s, '我们的思路', .85, 4.55, 4, .5, 20, TEAL, True)
    tb(s, '只用已经发生的访问历史，估计每个驻留页“下次还要多久才会被用到”，淘汰预计最晚被用到的页——在线地逼近 MIN。', .85, 5.05, 11.6, 1.2, 22)

    # 3 five algorithms -------------------------------------------------------
    s = slide('五种算法：它们靠什么选择受害页？',
              '前四个是课程要求的基本算法。FIFO 看进入内存的先后；LRU 看最近一次访问；Second chance 在 FIFO 基础上加访问位，给近期用过的页一次机会；MIN 看未来，是不可实现的理论下界。我们的算法同样只用历史，但同时看多种信号——间隔、前驱页、后继页，并由 25 个候选策略在线竞争，选当前最准的那个。',
              50)
    algs = [('FIFO', '进入内存最久', '公平，但不管页面是否仍在使用', '最简单', '9AA3B2'),
            ('LRU', '最近最久未访问', '假设最近用过的会再用', '循环访问超过内存时反而全错', AMBER),
            ('Second chance', 'FIFO + 访问位', '被访问过的页获得第二次机会', '只有 1 位信息，区分度低', BLUE),
            ('MIN (OPT)', '未来最晚才访问', '理论最优，现场不可用', '需要读取未来', SLATE),
            ('我们的算法', '预测下次访问时间', '间隔 + 上下文，多策略在线选择', '用计算换更少的缺页', TEAL)]
    w = 2.3
    for i, (n, sig, how, note, col) in enumerate(algs):
        x = .6 + i * (w + .15)
        box(s, x, 1.45, w, 4.9, 'FFFFFF', col)
        box(s, x, 1.45, w, .7, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=.25)
        tb(s, n, x, 1.5, w, .6, 19, 'FFFFFF', True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        tb(s, '依据', x + .1, 2.3, w - .2, .35, 13, GRAY)
        tb(s, sig, x + .1, 2.65, w - .2, .8, 18, col if col != '9AA3B2' else INK, True)
        tb(s, '做法', x + .1, 3.55, w - .2, .35, 13, GRAY)
        tb(s, how, x + .1, 3.9, w - .2, 1.0, 15)
        tb(s, '特点', x + .1, 4.95, w - .2, .35, 13, GRAY)
        tb(s, note, x + .1, 5.3, w - .2, 1.0, 15, INK)
    tb(s, '另外也实现了 Random、LFU 作为参照；以下对比以这五种为主。', .6, 6.5, 12, .45, 16, GRAY)

    # 4 frames charts -----------------------------------------------------------
    s = slide('不同页框数下的缺页次数',
              '横轴是页框数，纵轴是缺页次数，越低越好，虚线是 MIN 理论下界。观察：第一，FIFO、LRU、Second chance 在 5、10、20 页框时几乎重合在 100% 缺页——lu 和矩阵乘法是循环访问，循环长度超过内存时，这三个算法每次都换出马上要用的页。第二，我们的算法在所有页框下都贴近 MIN。第三，页框很大时，LRU 和 FIFO 才开始接近，但我们仍然更低。',
              60)
    for i, a in enumerate(FIVE):
        x = .6 + i * 2.5
        box(s, x, 1.28, .5, .12, COL[a], shape=MSO_SHAPE.RECTANGLE)
        tb(s, NAME[a] + (' 虚线' if a == 'opt' else ''), x + .55, 1.1, 1.95, .45, 14, INK, a == 'ours')
    pos = [(.45, 1.6), (6.75, 1.6), (.45, 4.3), (6.75, 4.3)]
    for (d, label), (x, y) in zip(TRACES, pos):
        line_chart(s, [str(f) for f in FRAMES], [(NAME[a], [val(d, f, a) for f in FRAMES], a) for a in FIVE], x, y, 6.2, 2.65, title=f'{label}（横轴：页框数；纵轴：缺页次数）', legend=False)

    # 5 core idea ----------------------------------------------------------------
    s = slide('核心思想：预测“下次访问还有多久”',
              'MIN 的做法是淘汰未来最晚被访问的页。我们没有未来，但可以估计：对每个驻留页，用它过去的访问间隔预测下一次访问的时间，再淘汰预测最远的那个。LRU 其实是一个特例——它假设最近访问的页下一次访问也最近。示意图中，A 的间隔很短、C 的间隔很长，所以淘汰 C。这个训练标签是页面真的再次出现之后才产生的，没有偷看未来。',
              45)
    # diagram: three pages with predicted next-use bars
    tb(s, '缺页发生，内存已满：谁最晚才会被再次用到？（示意）', .6, 1.3, 12, .5, 22, INK, True)
    rows = [('页面 A', .9, '预计很快再用', TEAL), ('页面 B', 2.6, '预计较晚再用', BLUE), ('页面 C', 6.1, '预计最晚再用  → 淘汰', 'C0392B')]
    for i, (n, ln, lab, col) in enumerate(rows):
        y = 2.1 + i * 1.0
        tb(s, n, .6, y, 1.6, .6, 20, INK, True, anchor=MSO_ANCHOR.MIDDLE)
        box(s, 2.2, y + .08, ln, .44, col, shape=MSO_SHAPE.RECTANGLE)
        tb(s, lab, 2.3 + ln, y, 6.5, .6, 18, col, True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, '← 现在                                                                    预测的下次访问时间 →', 2.2, 5.1, 10, .4, 13, GRAY)
    box(s, .6, 5.6, 12.1, 1.3, LIGHT)
    tb(s, '为什么比 LRU 好：LRU 只看“上次什么时候用过”；我们看“这页通常隔多久回来”。循环访问时，刚用过的页恰恰最晚才被再用。', .85, 5.7, 11.6, 1.1, 18)

    # 6 conditional reuse -----------------------------------------------------------
    s = slide('新信号：等待时间本身就是信息',
              '只用平均间隔会出问题。举一个示意例子：某页过去的访问间隔要么约 10，要么约 100。平均值是 55，哪一种情况都不是。如果它已经等了 20 还没出现，那么“间隔 10”的那一类已经排除，应当按 100 来预测。所以我们保存的是间隔分布，并且按当前已等待的时间做条件化。同时，同一页在不同程序路径下的间隔也不同，所以我们把前驱页、后继页作为上下文分开学习。参考了 LRB 的延迟标签、LHD 的“存活时间影响价值”以及 LeCaR 的在线组合思想，但都是自己实现、自己验证。',
              55)
    card(s, .6, 1.4, 5.9, 2.9, '只用平均值（容易错）',
         '示意：历史间隔 10、100、10、100\n平均 ≈ 55 —— 实际从不出现\n页面已等待 20，仍按 55 预测', head='9AA3B2', bsize=21)
    card(s, 6.8, 1.4, 5.9, 2.9, '条件复用分布（我们）',
         '已等待 20 → 排除“间隔 10”\n只保留大于 20 的历史间隔\n预测 ≈ 100', head=TEAL, bsize=21, fill='E3F4F2')
    tb(s, '怎样实现（有界内存）', .6, 4.5, 6, .5, 20, PURPLE, True)
    tb(s, '• 每组最多 16 个加权“间隔簇”，旧证据每次乘 0.98，适应访问模式变化\n• 分组上下文：最近间隔、前后间隔对、前驱页、后继页\n• 页面已被淘汰出历史表 → 记为“尚未返回”，加有限惩罚\n• “后继页”指已经到达的当前页，不是偷看下一页', .6, 5.0, 12.1, 2.0, 18, INK, space=5)

    # 7 architecture --------------------------------------------------------------------
    s = slide('怎么落地：25 个候选策略在线竞争',
              '实现上，每个请求只把当前页号交给算法。算法维护一份有界的历史观察；基于这份历史，25 个候选策略——包括 LRU、频率、片段重放、间隔预测、条件分布等——各自在“影子缓存”里模拟，影子里只有页号，不是真实页面。选择器用最近一段时间的衰减缺页损失给候选打分，领先超过门槛才切换，避免来回抖动。真实缓存永远只有 F 个页框，缺页时只淘汰一个页，并且只从真实驻留页里选，不被当前获胜者的影子保留。',
              50)
    stages = [('当前页', '每次只输入\n一个页号', 'A78BC9'), ('历史观察', '间隔分布\n前驱 / 后继', PURPLE), ('25 个候选策略', '影子缓存\n各自模拟', PURPLE),
              ('选择器', '衰减缺页损失\n+ 切换门槛', TEAL), ('真实缓存', 'F 个页框\n逐次合法淘汰', SLATE)]
    bw = 2.1
    for i, (t, sub, col) in enumerate(stages):
        x = .6 + i * (bw + .4)
        box(s, x, 2.0, bw, 2.0, col)
        tb(s, t, x, 2.15, bw, .6, 19, 'FFFFFF', True, PP_ALIGN.CENTER)
        tb(s, sub, x, 2.9, bw, 1.0, 16, 'FFFFFF', align=PP_ALIGN.CENTER)
        if i < 4:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + .05), Inches(2.75), Inches(.3), Inches(.5))
            a.fill.solid(); a.fill.fore_color.rgb = rgb(GRAY); a.line.fill.background()
    tb(s, '命中 / 缺页统计', 10.5, 4.15, 2.3, .4, 14, GRAY, align=PP_ALIGN.CENTER)
    card(s, .6, 4.7, 5.9, 2.1, '候选策略', '• LRU、频率、随机\n• 间隔预测、重放、条件分布、上下文 …', bsize=17)
    card(s, 6.8, 4.7, 5.9, 2.1, '选择规则', '• 损失 L ← (1 − 1/W)·L + 本次是否缺页\n• W = max(32, 8F)；领先超过 W/20 才切换', bsize=17, head=TEAL)

    # 8 rules compliance ---------------------------------------------------------------------
    s = slide('遵守规则：没有预取，也没有偷看未来',
              '这一页回答可能的质疑。程序逐行读入，对象只能看到当前页号，拿不到整个文件、也不知道长度；不会按文件名或页号大小选择算法。影子缓存只记页号，永远不会把影子里的页放进真实缓存。我们用 Release、ASan、UBSan 做了测试，并且在很多短序列上与穷举得到的精确 OPT 做了对照，还检查了前缀一致性、页号重命名不变、流式与一次性读入一致。',
              25)
    checks = [('只能 on-demand', '缺页才换入，每次只换入当前页，不预取'), ('不看未来', '接口一次只收到当前页号，无文件句柄、无长度'),
              ('真实页框不超限', '影子只含页号；真实缓存始终 ≤ F 页'), ('不依赖文件名 / 页号大小', '页号已是页号，不再除以页大小'),
              ('结果可复现', '随机算法使用 30 个固定种子；全部原始数据保留'), ('测试通过', 'Release + ASan + UBSan；与精确 OPT 对照')]
    for i, (t, d) in enumerate(checks):
        x = .6 + (i % 2) * 6.15; y = 1.4 + (i // 2) * 1.8
        box(s, x, y, 5.95, 1.55, 'FFFFFF', TEAL)
        tb(s, '✓', x + .15, y + .3, .7, .9, 38, TEAL, True)
        tb(s, t, x + .9, y + .15, 4.9, .5, 21, INK, True)
        tb(s, d, x + .9, y + .7, 4.9, .8, 16, GRAY)

    # 9 overall results ---------------------------------------------------------------------------
    s = slide('总体结果：四条课程 trace × 五档页框的平均缺页',
              f'这是最核心的结果：对四条课程 trace 和五档页框做等权平均。我们的算法平均缺页 {O["ours"]:,.0f}，比 LRU 少 {vs_lru:.1f}%（LRU 是三个基本算法里最好的），只比读取了未来的理论下界 MIN 多 {vs_opt:.1f}%。需要诚实说明：这四条 trace 在开发时用过，所以它们不是盲测，后面会用冻结后新生成的序列验证。',
              50)
    order = ['opt', 'ours', 'random', 'lru', 'second_chance', 'fifo', 'lfu']
    bar_chart(s, [NAME[a] for a in order], [round(O[a]) for a in order], [COL[a] if a in FIVE else 'C9CED8' for a in order], .5, 1.2, 8.6, 5.7)
    box(s, 9.4, 1.4, 3.4, 5.3, LIGHT)
    tb(s, f'−{vs_lru:.1f}%', 9.5, 1.6, 3.2, .9, 44, TEAL, True, PP_ALIGN.CENTER)
    tb(s, '相对 LRU 的缺页减少', 9.5, 2.5, 3.2, .4, 15, GRAY, align=PP_ALIGN.CENTER)
    tb(s, f'+{vs_opt:.1f}%', 9.5, 3.5, 3.2, .9, 44, SLATE, True, PP_ALIGN.CENTER)
    tb(s, '距离 MIN 理论下界', 9.5, 4.4, 3.2, .4, 15, GRAY, align=PP_ALIGN.CENTER)
    tb(s, f'基本算法约 2.4–2.5 万\n我们 {O["ours"]:,.0f}', 9.5, 5.2, 3.2, 1.0, 16, INK, align=PP_ALIGN.CENTER)

    # 10 per trace table -------------------------------------------------------------------------------
    s = slide('不同访问序列：每条 trace 的平均缺页',
              '每一行是一条 trace 在五档页框上的平均缺页次数。我们的算法在 lu、mm16、mm32 上明显优于三个经典算法；在 qsort 上，因为 qsort 只有 524 次访问、192 个不同页，没有太多可学习的重复，所以各算法差距小，我们依然更低。绿色列是我们的算法。',
              30)
    rows = [[lab, *[f'{tmean(d, a):,.1f}' for a in ['fifo', 'lru', 'second_chance', 'opt', 'ours']], f'{100 * (1 - tmean(d, "ours") / min(tmean(d, "fifo"), tmean(d, "lru"), tmean(d, "second_chance"))):.1f}%'] for d, lab in TRACES]
    table(s, ['Trace', 'FIFO', 'LRU', 'Second chance', 'MIN (OPT)', '我们的算法', '较最佳经典算法少'], rows, y=1.5, size=18, hl_col=5, rowh=.75)
    tb(s, '数字为五档页框（5、10、20、50、100）缺页次数的平均值。', .6, 5.5, 12, .4, 16, GRAY)
    tb(s, '完整的逐页框数字见附录。', .6, 5.95, 12, .4, 16, GRAY)

    # 11 parameters ------------------------------------------------------------------------------------------
    sweep = list(csv.DictReader((Path(__file__).parent / 'param_sweep.csv').open()))
    s = slide('算法参数的影响',
              '算法有两个主要参数。窗口系数决定选择器比较候选策略近期表现的时间窗口，历史系数决定每个页面可以保留多少历史。图中是四条课程 trace 平均缺页，越低越好。结论：窗口系数在 2 到 8 都比较稳定，太小则选择器过于敏感；历史系数太小时记不住足够的间隔信息，缺页明显上升，16 附近最好，继续加大收益消失。我们使用窗口系数 4、历史系数 16。这两个参数是在课程和开发数据上选的，所以不能当作对未知数据的保证。',
              40)
    for k, (kind, label, x) in enumerate([('window', '窗口系数（选择器比较窗口）', .45), ('history', '历史系数（每页历史容量）', 6.75)]):
        vs = ['1', '2', '4', '8', '16', '32']
        avg = []
        for v in vs:
            avg.append(statistics.mean(float(r['avg_faults']) for r in sweep if r['kind'] == kind and r['value'] == v))
        c = bar_chart(s, vs, [round(a) for a in avg], [TEAL if v in ('4', '16') and ((kind == 'window' and v == '4') or (kind == 'history' and v == '16')) else '9AA3B2' for v in vs], x, 1.7, 6.2, 4.2, size=13)
        tb(s, label, x, 1.15, 6.2, .5, 20, PURPLE, True, PP_ALIGN.CENTER)
        c.value_axis.minimum_scale = round(min(avg) * .9, -2)
    tb(s, '纵轴：四条课程 trace 的平均缺页次数（已截断起点）；绿色为我们采用的取值。', .6, 6.05, 12.1, .4, 15, GRAY)
    tb(s, '默认：窗口系数 = 4，历史系数 = 16。', .6, 6.5, 12.1, .4, 17, INK, True)

    # 12 holdout ---------------------------------------------------------------------------------------------------
    sets = []
    for fn, lab in [('development', '开发用 10 类'), ('regression', '回归 8 类'), ('holdout', '冻结后新生成 10 类')]:
        m = read(fn + '_means')
        sets.append((lab, {a: statistics.mean(float(r['average_faults']) for r in m if r['algorithm'] == key(a)) for a in FIVE}))
    s = slide('换一批序列：构造的循环、热点、Zipf、混合程序',
              '老师现场的 trace 不在给定数据集里，所以我们还生成了很多访问模式不同的序列：循环、矩阵、热点加扫描、Zipf、多个程序交错等。最右边一组是在算法冻结之后才生成、完全没有用来调参的 10 类序列。我们的算法仍然最少，比 LRU 少约 15%，距离 MIN 约 12%。要说明的是：也有不占优的：在 matrix_ikj_24 上比 LRU 多约 1%，在完全随机均匀访问 uniform_997 上与 LRU 持平——随机访问没有规律可学，任何算法都帮不上忙。这些都如实记录在报告里。',
              45)
    cats = [n for n, _ in sets]
    data = CategoryChartData(); data.categories = cats
    for a in FIVE: data.add_series(NAME[a], [round(d[a]) for _, d in sets])
    c = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.5), Inches(1.2), Inches(8.8), Inches(5.7), data).chart
    c.has_legend = True; c.legend.position = XL_LEGEND_POSITION.BOTTOM; c.legend.include_in_layout = False; c.legend.font.size = Pt(13); c.legend.font.name = FONT
    style_axis(c, 13); c.plots[0].gap_width = 60; c.plots[0].overlap = -5
    for ser, a in zip(c.series, FIVE):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(COL[a])
    last = sets[2][1]
    box(s, 9.5, 1.4, 3.3, 5.3, LIGHT)
    tb(s, '新构造序列', 9.6, 1.5, 3.1, .45, 18, PURPLE, True, PP_ALIGN.CENTER)
    tb(s, f'−{100 * (1 - last["ours"] / last["lru"]):.1f}%', 9.6, 2.1, 3.1, .9, 42, TEAL, True, PP_ALIGN.CENTER)
    tb(s, '相对 LRU', 9.6, 3.0, 3.1, .4, 15, GRAY, align=PP_ALIGN.CENTER)
    tb(s, f'+{100 * (last["ours"] / last["opt"] - 1):.1f}%', 9.6, 3.7, 3.1, .9, 42, SLATE, True, PP_ALIGN.CENTER)
    tb(s, '距离 MIN', 9.6, 4.6, 3.1, .4, 15, GRAY, align=PP_ALIGN.CENTER)
    tb(s, '这 10 类序列在算法冻结后才生成，未用于调参；个别序列略差于 LRU（如 matrix_ikj_24），已如实记录。', 9.6, 5.15, 3.1, 1.4, 13, GRAY)

    # 13 cost + limits ---------------------------------------------------------------------------------------------------
    s = slide('开销与局限',
              f'代价方面：25 个影子缓存让模拟时间变长，mm32 五档合计约 {perf["median_milliseconds"]["5"] / 1000:.1f} 秒，这是本机上程序端到端的模拟时间，不是操作系统真实的缺页延迟。内存是有界的，一百万行的纯扫描输入峰值约 20 MiB。局限方面：它依赖访问模式有重复；完全随机或第一次出现的页面无法预测；我们也不能保证在未知文件上拿第一。',
              30)
    card(s, .6, 1.4, 5.9, 2.7, '代价', f'• mm32（67,235 次访问）五档合计约 {perf["median_milliseconds"]["5"] / 1000:.1f} 秒\n• 一百万行纯扫描：峰值内存约 {perf["unique_scan_scaling"][-1]["peak_rss_kib"] / 1024:.0f} MiB\n• 内存随页框数有界，与输入长度无关', bsize=19)
    card(s, 6.8, 1.4, 5.9, 2.7, '局限', '• 依赖访问模式有重复\n• 完全随机、第一次出现的页无法预测\n• 不保证任意未知输入都第一', bsize=19, head='C0392B')
    box(s, .6, 4.45, 12.1, 2.35, 'E3F4F2')
    tb(s, '现场使用', .85, 4.55, 4, .5, 20, TEAL, True)
    tb(s, '直接读取老师给的文件，一行一个页号，输出五档缺页次数和平均值。\n演练（mm32）：' + ' · '.join(f'{x:,.0f}' for x in mm) + f'，平均 {statistics.mean(mm):,.2f}', .85, 5.1, 11.6, 1.5, 19)

    # 14 summary ----------------------------------------------------------------------------------------------------------------------
    s = slide('总结',
              '总结三点：第一，实现了 FIFO、LRU、Second chance、MIN 四个基本算法和我们自己的自适应算法，并做了页框数、不同 trace、不同参数的实验。第二，核心贡献是用条件复用分布预测下次访问时间，并让 25 个候选策略在线竞争。第三，在课程 trace 上平均缺页比 LRU 少三成多，只比理论下界多约两个百分点；在冻结后新生成的序列上仍然领先；同时我们也说明了局限。谢谢。',
              20)
    pts = [('1', '五种算法完整对比', 'FIFO · LRU · Second chance · MIN · 我们的算法；5 种页框、4 条 trace、多组参数'),
           ('2', '核心贡献', '条件复用分布预测下次访问时间；25 个候选策略在线竞争，真实页框始终合法'),
           ('3', '结果', f'课程 trace 平均缺页比 LRU 少 {vs_lru:.0f}%，仅比理论下界 MIN 多 {vs_opt:.1f}%；新构造序列仍领先'),
           ('4', '诚实说明', '依赖访问模式重复；未知文件不保证第一；四条课程 trace 参与过开发')]
    for i, (n, t, d) in enumerate(pts):
        y = 1.4 + i * 1.35
        box(s, .6, y, 12.1, 1.15, LIGHT)
        c = box(s, .8, y + .2, .75, .75, TEAL, shape=MSO_SHAPE.OVAL)
        tb(s, n, .8, y + .2, .75, .75, 26, 'FFFFFF', True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        tb(s, t, 1.8, y + .1, 10.7, .5, 21, INK, True)
        tb(s, d, 1.8, y + .58, 10.7, .5, 16, GRAY)

    # appendix -----------------------------------------------------------------------------------------------------------------------------------------
    keys = ['random', 'fifo', 'opt', 'lfu', 'lru', 'second_chance', 'ours']
    for d, lab in TRACES:
        s = slide(f'附录：{lab} 完整缺页次数', '附录，不讲。Random 为 30 个固定种子的均值，其余均为确定性整数。', 0)
        table(s, ['页框数', 'Random', 'FIFO', 'MIN (OPT)', 'LFU', 'LRU', 'Second chance', '我们的算法'],
              [[f, *[f'{val(d, f, a):,.1f}'.rstrip('0').rstrip('.') for a in keys]] for f in FRAMES], y=1.6, size=16, hl_col=7, rowh=.7)

    prs.save(ROOT / 'presentation/contest-v5.pptx')
    text = ['# 汇报讲稿', f'共 {len(prs.slides)} 页，前 14 页约 10 分钟，其余为附录。']
    for i, (title, sec, script) in enumerate(notes, 1): text += [f'## {i}. {title}（{sec} 秒）', script]
    (ROOT / 'presentation/V5_SPEAKER_NOTES.md').write_text('\n\n'.join(text) + '\n')
    print('slides', len(prs.slides), 'total seconds', sum(n[1] for n in notes[:14]))


if __name__ == '__main__': main()
