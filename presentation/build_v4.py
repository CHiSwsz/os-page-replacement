#!/usr/bin/env python3
"""Build the V4 classroom deck without an external template or cloud service."""
import csv
import json
from pathlib import Path
import statistics

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/v4'
PURPLE='6536A0'; INK='222638'; GRAY='637083'; TEAL='087F78'
FONT='Droid Sans Fallback'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
notes=[]


def read(name):
    with (OUT/(name+'.csv')).open() as f:return list(csv.DictReader(f))


def textbox(slide,text,x,y,w,h,size=22,color=INK,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.04)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(10)
    return shape


def slide(title,script,source='REPORT.md',seconds=40):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(7.12),prs.slide_width,Inches(.38))
    bar.fill.solid();bar.fill.fore_color.rgb=RGBColor.from_string('F0EAF8');bar.line.fill.background()
    textbox(s,title,.52,.28,12.3,.7,30,PURPLE,True)
    textbox(s,'Adaptive V4  ·  Online demand paging',.55,7.16,10,.24,10,GRAY)
    textbox(s,str(len(prs.slides)),12.15,7.15,.6,.24,10,GRAY)
    s.notes_slide.notes_text_frame.text=f'建议用时：{seconds} 秒\n{script}\n数据与实现来源：{source}'
    notes.append((title,seconds,script,source))
    return s


def bullets(s,lines,y=1.35,size=23):
    textbox(s,'\n'.join('• '+line for line in lines),.72,y,11.9,5.55-y,size)


def table(s,headers,rows,y=1.5,size=16):
    shape=s.shapes.add_table(len(rows)+1,len(headers),Inches(.6),Inches(y),Inches(12.1),Inches(min(4.9,.58*(len(rows)+1))))
    for i,row in enumerate([headers,*rows]):
        for j,txt in enumerate(row):
            cell=shape.table.cell(i,j);cell.text=str(txt)
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(PURPLE if i==0 else ('F6F3FA' if i%2 else 'FFFFFF'))
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(size);p.font.bold=(i==0);p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else INK)
    return shape


def chart(s,categories,series,x,y,w,h,kind=XL_CHART_TYPE.COLUMN_CLUSTERED):
    data=CategoryChartData();data.categories=categories
    for name,values in series:data.add_series(name,values)
    c=s.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    c.has_legend=len(series)>1
    if c.has_legend:
        c.legend.position=XL_LEGEND_POSITION.BOTTOM;c.legend.include_in_layout=False;c.legend.font.size=Pt(11)
    c.category_axis.tick_labels.font.size=Pt(11);c.value_axis.tick_labels.font.size=Pt(10)
    c.value_axis.minimum_scale=0
    c.value_axis.has_major_gridlines=True
    for i,ser in enumerate(c.series):
        color=['94A3B8','F59E0B',TEAL,'334155'][i%4]
        if kind==XL_CHART_TYPE.LINE_MARKERS:
            ser.format.line.color.rgb=RGBColor.from_string(color);ser.format.line.width=Pt(2)
        else:
            ser.format.fill.solid();ser.format.fill.fore_color.rgb=RGBColor.from_string(color)
    return c


def main():
    summary=read('course_summary');means=read('course_means');holdout=read('holdout_means')
    overall={a:statistics.mean(float(r['average_faults']) for r in means if r['algorithm']==a) for a in dict.fromkeys(r['algorithm'] for r in means)}
    val=lambda d,f,a:float(next(r['faults_mean'] for r in summary if r['dataset']==d and int(r['frames'])==f and r['algorithm']==a))
    s=slide('让历史预测更有用：Adaptive V4','今天介绍一个只看过去的页面替换组合策略。我们在队友 V3 的基础上继续优化，重点是现场能读未知文件并按规定报告五档平均。OPT 只是参照，正式算法不看未来。',seconds=20)
    textbox(s,'更低的缺页次数，完整的在线约束',.78,1.65,11.7,.8,32,INK,True)
    textbox(s,f"课程平均缺页  {overall['adaptive']:,.2f}\n较 LRU 减少  {100*(1-overall['adaptive']/overall['lru']):.2f}%\n较 V3 进一步减少  {100*(1-overall['adaptive']/overall['adaptive_v3']):.3f}%",.85,2.85,11.5,2.6,27,TEAL)
    textbox(s,'4 条课程 trace × 5 档页框 · seed 0–29 · 所有原始数据可复现',.85,6.25,11.6,.5,18,GRAY)
    s=slide('任务与评分：只比较实际缺页', '五档各自从空内存开始。一次访问只能命中或者缺页，只有缺页才装入当前页。报告最后一行是五个整数缺页次数相加除以五，不是缺页率平均。',seconds=35)
    bullets(s,['输入：每行一个页号；课程数据为十六进制。','页框：5、10、20、50、100；每档独立空内存启动。','缺页且已满：淘汰一页，只装入当前请求页。','评分：五档缺页次数的算术平均，越少越好。','现场入口逐行读入，不把未来序列交给算法。'])
    s=slide('对照方法与 V3 的起点','FIFO、MIN、LRU、Second chance 是要求的四个对照；另外保留 Random 和 LFU。LFU 是当前驻留期间的频次，卸载后重置。V3 已经通过历史复用间隔取得很好的效果，因此新增收益不是从弱基线开始。',seconds=35)
    table(s,['方法','规则'],[['FIFO','淘汰最早装入的页'],['LRU / Second chance','最近使用顺序 / 环形访问位'],['Random / LFU','均匀随机 / 驻留期频次，平局按 LRU'],['MIN / OPT','读取真实未来，仅作离线下界'],['Adaptive V3','Random + LRU + 历史间隔预测']],size=18)
    s=slide('七个专家：热页、周期与新页','Gap 是每个页面的间隔指数平滑。Replay 匹配过去的相同访问后缀。Frequency 应对稳定热点。新页没有自己的间隔，所以 Cold prior 借用近期复用间隔的统计；Signature 从相同上下文其他页学到间隔。两个新页机制都可能出错，由在线选择器约束。',source='src/context.cpp；REPORT.md §2.1',seconds=55)
    table(s,['专家','信息'],[['Random / LRU','随机与最近使用，基础对照'],['Gap / Replay','单页复用间隔 / 已发生的相似片段'],['Frequency','随时间衰减的频率'],['Cold prior','最近 64 个已完成复用间隔的下四分位'],['Signature','锚点页号 + 相对位置，实际返回后才训练']],size=18)
    s=slide('衰减损失 + 切换门槛，抑制振荡','七个影子都只处理历史。每次完成请求，把该策略的缺页记成一，命中记成零，累积时衰减旧证据。选专家前只能读上一次的损失。周期检查并要求显著超过门槛，避免循环中一点波动就切换。',source='src/context.cpp: OnlineAdaptive::Impl',seconds=60)
    textbox(s,'Lⱼ ← (1 − 1/W) Lⱼ + 本次缺页指示',.85,1.55,11.8,.85,30,PURPLE,True)
    bullets(s,['默认 W = 2 × max(32, 4F)。','完成 W 次请求后，每 W/4 次检查一次。','新专家必须少超过 max(1, floor(W/20)) 才切换。','F ≤ 20 从 Gap 起步，其余从 LRU 起步。','规则按容量统一设定，不看文件名，也不挑现场种子。'],y=2.75,size=22)
    s=slide('影子只有元数据；真实缓存仍是 F 页','这是最重要的合法性说明。当前页进入影子后，真实缓存若缺页，淘汰它有但影子不保留的某页。只装入当前请求页，逐步接近所选策略。真实命中时不为同步影子而替换。影子模拟开销必须承认，但不能把页号元数据当作多缓存了页面内容。',seconds=50)
    for x,label in [(.8,'当前页号'),(4.45,'七个影子：页号/统计'),(8.6,'真实 F 个物理页框')]:
        box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(1.75),Inches(3.1),Inches(1.1))
        box.fill.solid();box.fill.fore_color.rgb=RGBColor.from_string('F0EAF8');box.line.color.rgb=RGBColor.from_string(PURPLE)
        textbox(s,label,x+.15,2.04,2.85,.5,19,PURPLE,True)
    textbox(s,'→',3.95,1.95,.5,.5,28);textbox(s,'→',7.9,1.95,.5,.5,28)
    bullets(s,['专家切换不复制影子驻留集合。','真实缺页时，从合法驻留页里选一个淘汰对象。','装入的始终只有当前请求页；命中不移动、不替换。'],y=3.55,size=24)
    s=slide('课程总体：比 V3 再降 0.576%','这是四条轨迹乘五档页框的等权平均。Random、V3、V4 各使用相同的三十个种子。V4 在二十个配置上十五个更好、五个持平，但它仍比离线最优高约百分之四点二五。',source='results/v4/course_summary.csv',seconds=45)
    names=['OPT','V4','V3','Random','LRU','Clock','FIFO','LFU'];keys=['opt','adaptive','adaptive_v3','random','lru','second_chance','fifo','lfu']
    c=chart(s,names,[('Mean page faults',[overall[k] for k in keys])],.75,1.35,11.9,4.8)
    for i,p in enumerate(c.series[0].points):p.format.fill.solid();p.format.fill.fore_color.rgb=RGBColor.from_string(TEAL if i==1 else PURPLE if i==2 else 'CBD5E1')
    textbox(s,'15 个配置改善 · 5 个持平 · 0 个退化（相对 V3 均值）',.85,6.25,11.9,.5,20,TEAL,True)
    s=slide('分轨迹观察：收益取决于局部结构','小页框最容易让 LRU 抖动。mm32 的五十页框从二万一千一百七十九降至一万九千九百零九，新上下文预测起了作用。mm16 的一百页框也改善。qsort 一百页框仍不如 Random 的平均表现，不能只报有利配置。',source='results/v4/course_summary.csv',seconds=60)
    s.shapes.add_picture(str(OUT/'course_comparison.png'),Inches(.8),Inches(1.22),width=Inches(11.65),height=Inches(5.5))
    s=slide('参数实验：默认并非每组的最优','窗口因子二在公开课程数据上的平均更低，窗口八反应更慢。历史太短容易遗忘，太长也可能留下旧模式。默认四和十六在额外验证前固定，现场不根据全文件试出最好参数。',source='results/v4/parameters_summary.csv',seconds=35)
    s.shapes.add_picture(str(OUT/'parameters.png'),Inches(.8),Inches(1.35),width=Inches(11.65),height=Inches(4.9))
    textbox(s,'3,000 条参数运行记录；收益与切换稳定性存在取舍。',.9,6.36,11.5,.42,18,GRAY)
    s=slide('冻结后 8 类构造测试：保留不利结果','新增八类数据是在规则固定后生成并运行的，不能叫老师的隐藏测试。总体约减少百分之一点一一。带噪循环小幅退化，Markov 分块仍略差于 LRU；随机数据上很小的数字差不代表真的能预测随机序列。',source='results/v4/holdout_means.csv',seconds=50)
    s.shapes.add_picture(str(OUT/'holdout_comparison.png'),Inches(.8),Inches(1.35),width=Inches(11.65),height=Inches(5.3))
    perf=json.loads((OUT/'performance.json').read_text());large=perf['unique_scan_scaling'][-1]
    s=slide('开销与验证：省缺页，不是省模拟时间','维护七个影子不是免费开销。我们说明真实缓存和元数据的区别，也报告实测运行时间。百万行流式输入正确完成，数百万项断言及内存和未定义行为检查通过。Windows 入口已提供，但这里的实机验证是 Linux。',source='results/v4/performance.json；tests/',seconds=45)
    bullets(s,[f"mm32 五档运行：V3 {perf['median_milliseconds']['3']:.0f} ms，V4 {perf['median_milliseconds']['4']:.0f} ms（本机中位数）。",f"百万次纯扫描：平均缺页 1,000,000，耗时 {large['elapsed_seconds']:.2f} 秒。",'算法元数据 O(F)，时间期望 O(NF)；token 字符串映射另占 O(U)。','250 万余项断言；Release、ASan、UBSan 全部通过。','逐步检查容量、事件、OPT 下界、未来前缀不变与页号重命名。'],size=21)
    s=slide('现场只换文件路径','提前编译并用 mm32 演练，结果应和这一页相同。现场只替换文件路径，不读完整文件后选择策略，不反复挑种子。若老师文件是十进制或字符串标签，用对应格式选项，含义在快速说明中写清楚。',source='QUICKSTART.md',seconds=30)
    textbox(s,'bash run.sh "老师给的文件.txt"',.9,1.5,11.6,.7,27,PURPLE,True)
    textbox(s,'5 / 10 / 20 / 50 / 100 页框\n自动输出每档缺页数、缺页率和最后的平均缺页次数',.95,2.65,11.5,1.3,23)
    textbox(s,'mm32 演练：55,006 · 49,256 · 39,016 · 19,909 · 19,604\nAverage page faults (5 capacities): 36558.20',.95,4.45,11.5,1.25,20,TEAL)
    s=slide('可以保证的，和仍有局限的','我们能保证实验过程可复现、在线接口不接收未来数据、真实缓存不超过物理页框数，并如实保留所有结果。我们不能保证对任意未知轨迹第一；优势来自对模式的学习，代价是元数据、计算和切换滞后。谢谢。',seconds=20)
    bullets(s,['保证：真实 demand paging、固定规则、完整原始记录。','实测：课程相对 V3 再降低 0.576%，相对 LRU 降低 30.43%。','代价：更多元数据、计算和学习时间。','局限：预测可能失准；未知轨迹没有必胜保证。','代码、说明、数据、报告、讲稿和 PPT 均可本地复现。'])
    keys=['random','fifo','opt','lfu','lru','second_chance','adaptive_v3','adaptive']
    for ds in ['qsort','lu','mm16','mm32']:
        s=slide('附录：'+ds+' 完整缺页计数','所有计数直接来自 course_summary.csv。随机方法为三十个种子的均值；V4 这四条轨迹三十个种子的结果相同。',source='results/v4/course_summary.csv',seconds=0)
        rows=[[f,*[f'{val(ds,f,a):,.2f}' if a in ['random','adaptive_v3'] else f'{val(ds,f,a):,.0f}' for a in keys]] for f in [5,10,20,50,100]]
        table(s,['Frames','Random','FIFO','OPT','LFU','LRU','Clock','V3','V4'],rows,y=1.7,size=13)
        textbox(s,'来源：30 个固定种子，原始整数与校验值保存在 results/v4/。',.8,6.2,11.8,.5,18,GRAY)
    target=ROOT/'presentation/contest-v4.pptx';prs.save(target)
    text=['# V4 汇报讲稿与答辩','前 13 页约 9 分钟，留 1 分钟过渡；后 4 页为数据附录。']
    for i,(title,seconds,script,source) in enumerate(notes,1):
        text += [f'## {i}. {title}（{seconds} 秒）',script,'数据与实现来源：`'+source+'`']
    text += ['## 可能被问到的问题',
        '**影子是不是偷偷增加内存？** 影子保存页号和统计，不保存页面内容。我们明确报告 O(F) 元数据开销；逐步事件验证真实缓存最多 F 页。',
        '**读完整文本再做 OPT 吗？** 现场入口逐行调用 access(page)，API 没有未来数组或文件句柄。OPT 只存在于对照实验入口。',
        '**为什么新页也能预测？** 先验来自其他页面已经完成的复用间隔；签名来自过去相同锚点与偏移上下文。预测可能错误，所以需要在线比较。',
        '**为什么不用参数实验中课程平均更好的窗口 2？** 默认规则在额外验证前固定，考虑整个开发集和切换稳定性；参数研究说明敏感性，不保证某个参数对未知输入最优。',
        '**能否保证第一？** 不能。未知轨迹可能没有可预测结构，甚至专门对抗我们的规则；报告给出了噪声循环退化和不如 LRU 的情况。',
        '**复杂度？** 固定七个影子，淘汰与片段预测扫描 O(F)，N 次访问期望 O(NF)；数值页号流式模式元数据 O(F)，字符串模式额外 O(U)。',
        '**来源与学术创新？** 参考在线策略学习和序列预测思想，在本作业中作工程组合，没有声称继承论文竞争比或发明全新理论。']
    (ROOT/'presentation/V4_SPEAKER_NOTES.md').write_text('\n\n'.join(text)+'\n',encoding='utf-8')
    print('Saved',target,'with',len(prs.slides),'slides and speaker notes')


if __name__=='__main__':main()
